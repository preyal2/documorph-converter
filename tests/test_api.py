"""
Integration tests for FastAPI REST API endpoints using TestClient.
"""

import sys
import unittest
from pathlib import Path
from docx import Document
from starlette.testclient import TestClient

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from backend.app import app

class TestDocuMorphAPI(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.fixtures_dir = PROJECT_ROOT / "tests" / "fixtures"
        cls.fixtures_dir.mkdir(parents=True, exist_ok=True)
        cls.docx_path = cls.fixtures_dir / "api_test.docx"

        # Create a sample docx
        doc = Document()
        doc.add_heading("API Test Title", level=0)
        doc.add_paragraph("Testing FastAPI file upload and conversion pipeline.")
        doc.save(str(cls.docx_path))

    @classmethod
    def tearDownClass(cls):
        if cls.docx_path.exists():
            try:
                cls.docx_path.unlink()
            except Exception:
                pass

    def test_health_check(self):
        """Test GET /api/health."""
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "healthy")
        self.assertTrue(data["capabilities"]["word_to_pdf"])
        self.assertTrue(data["capabilities"]["pdf_to_word"])

    def test_convert_word_to_pdf_endpoint(self):
        """Test POST /api/convert/word-to-pdf."""
        with open(self.docx_path, "rb") as f:
            response = self.client.post(
                "/api/convert/word-to-pdf",
                files={"file": ("api_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertTrue("download_url" in data)
        self.assertTrue(data["output_filename"].endswith(".pdf"))

        # Test download endpoint
        dl_response = self.client.get(data["download_url"])
        self.assertEqual(dl_response.status_code, 200)
        self.assertGreater(len(dl_response.content), 0)

    def test_invalid_extension(self):
        """Test that invalid file extensions return 400."""
        response = self.client.post(
            "/api/convert/word-to-pdf",
            files={"file": ("malicious.exe", b"invalid payload", "application/octet-stream")}
        )
        self.assertEqual(response.status_code, 400)

if __name__ == "__main__":
    unittest.main()
